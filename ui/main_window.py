# ui/main_window.py

import dearpygui.dearpygui as dpg
from core import app_state
from utils import api_client, template_fetcher, upload_manager
import pandas as pd
import json
from openpyxl.worksheet.datavalidation import DataValidation
from . import misc_tools_window  # Import the new misc_tools_window

# --- Module-level Global Variables ---
_log_messages = []
_selected_file_path = ""  # To store the path of the selected Excel file

# NEW: Module-level global to store config, accessible by module-level callbacks
_main_window_config_global = None


# NEW: log_message moved to module level and is now directly callable
def log_message(message):
    """A function to send messages to the log window."""
    print(message)
    _log_messages.append(str(message))
    if dpg.does_item_exist("log_text"):  # Check if DPG item exists before updating
        dpg.set_value("log_text", "\n".join(_log_messages))
    if dpg.does_item_exist("log_window"):
        dpg.set_y_scroll("log_window", -1.0)


# NEW: fetch_hubs_callback moved to module level
def fetch_hubs_callback():
    global _main_window_config_global  # Access the global config object

    if _main_window_config_global is None:
        log_message("❌ Configuration not loaded. Cannot fetch data.")
        return

    force_fetch = dpg.get_value("force_refresh_checkbox")

    _log_messages.clear()
    dpg.configure_item("status_text", default_value="Fetching data...")
    dpg.configure_item("fetch_button", enabled=False)

    log_message("Fetching all documents...")
    from utils.token_checker import refresh_access_token
    # Use _main_window_config_global for API calls
    app_state.set_all_documents(
        api_client.get_all_documents(_main_window_config_global, log_callback=log_message, force_api_fetch=force_fetch))

    if not app_state.all_documents:
        log_message("❌ No documents found. Ensure your token is valid and URL is correct.")
        dpg.configure_item("fetch_button", enabled=True)
        dpg.configure_item("status_text", default_value="Failed to fetch documents.")
        return

    log_message("🔍 Fetching master list of all templates (with details)...")
    # Use _main_window_config_global for API calls
    app_state.set_all_available_templates(
        api_client.get_all_templates(_main_window_config_global, log_callback=log_message, force_api_fetch=force_fetch))
    if not app_state.all_available_templates:
        log_message("❌ No templates found. This might affect template selection.")
        dpg.configure_item("status_text", default_value="No templates found.")

    app_state.hub_names_for_combo = sorted(
        list(set(
            doc.get('document_hub_id') for doc in app_state.all_documents if doc.get('document_hub_id') is not None))
    )

    hub_display_names = []
    for h_id in app_state.hub_names_for_combo:
        hub_doc = next(
            (doc for doc in app_state.all_documents if doc.get('id') == h_id and doc.get('parent_folder_id') is None),
            None)
        if hub_doc and hub_doc.get('title'):
            hub_display_names.append(f"{hub_doc['title']} (ID: {h_id})")
        else:
            hub_display_names.append(f"Hub ID: {h_id}")

    if hub_display_names:
        dpg.configure_item("hub_combo", items=hub_display_names, enabled=True)
        log_message(f"✅ Found {len(hub_display_names)} unique Document Hubs.")
        dpg.configure_item("status_text", default_value="✅ Select a Document Hub.")
    else:
        dpg.configure_item("hub_combo", items=[], enabled=False)
        log_message("❌ No Document Hub IDs found in the fetched documents.")
        dpg.configure_item("status_text", default_value="No Document Hubs available.")

    dpg.configure_item("fetch_button", enabled=True)


# ADD THIS NEW HELPER FUNCTION (if you're using the menu bar item)
def _show_and_refresh_misc_tools():
    """Helper function to show the misc tools window and refresh its data."""
    dpg.configure_item("misc_tools_window", show=True)
    misc_tools_window.refresh_misc_data_and_combos(None, None)


def create_main_window(config: dict, status_message: str):
    global _main_window_config_global
    _main_window_config_global = config  # Store config at module level

    def show_config():
        """Switches view to the configuration window."""
        dpg.configure_item("main_window", show=False)
        dpg.configure_item("config_window", show=True)
        dpg.set_primary_window("config_window", True)

    def hub_selected_callback():
        selected_hub_value = dpg.get_value("hub_combo")
        if not selected_hub_value: return

        if " (ID: " in selected_hub_value:
            hub_id_str = selected_hub_value.split(" (ID: ")[1].rstrip(")")
        else:
            hub_id_str = selected_hub_value.replace("Hub ID: ", "")

        app_state.set_selected_hub(int(hub_id_str), selected_hub_value)

        log_message(f"\n▶️ Hub '{selected_hub_value}' selected. ID: {app_state.selected_hub_id}")

        dpg.configure_item("folder_combo", items=[], enabled=False, default_value="")
        dpg.configure_item("template_combo", items=[], enabled=False, default_value="")

        if dpg.does_item_exist("generate_template_button"):
            dpg.configure_item("generate_template_button", enabled=False)
        if dpg.does_item_exist("upload_button"):
            dpg.configure_item("upload_button", enabled=False)

        app_state.set_selected_folder(None, "")
        app_state.set_selected_template(None, None)

        log_message(f"🔍 Fetching folders for Document Hub ID: {app_state.selected_hub_id}...")
        _folders_data_local = template_fetcher.get_folders_for_hub(_main_window_config_global,
                                                                   app_state.selected_hub_id, log_callback=log_message)

        app_state.folder_titles_for_combo = []

        hub_doc = next((doc for doc in app_state.all_documents if
                        doc.get('id') == app_state.selected_hub_id and doc.get('parent_folder_id') is None), None)
        if hub_doc and hub_doc.get('title'):
            app_state.folder_titles_for_combo.append(f"{hub_doc['title']} (Root of Hub)")
        else:
            app_state.folder_titles_for_combo.append(f"Root of Hub (ID: {app_state.selected_hub_id})")

        for folder in _folders_data_local:
            if folder.get('title') and folder.get('id'):
                app_state.folder_titles_for_combo.append(f"{folder['title']} (Folder ID: {folder['id']})")

        if app_state.folder_titles_for_combo:
            dpg.configure_item("folder_combo", items=app_state.folder_titles_for_combo, enabled=True)
            dpg.set_value("folder_combo", app_state.folder_titles_for_combo[0])
            log_message(
                f"✅ Found {len(app_state.folder_titles_for_combo)} selectable folders/root for Hub {app_state.selected_hub_id}.")
            folder_selected_callback()
        else:
            dpg.configure_item("folder_combo", items=[f"Root of Hub (ID: {app_state.selected_hub_id})"], enabled=True)
            dpg.set_value("folder_combo", f"Root of Hub (ID: {app_state.selected_hub_id})")
            log_message(
                f"❌ No sub-folders found for Hub {app_state.selected_hub_id}. You can still upload directly to the hub root.")
            folder_selected_callback()

        template_ids_in_selected_hub_documents = set(doc.get('template_id') for doc in app_state.all_documents
                                                     if doc.get(
            'document_hub_id') == app_state.selected_hub_id and doc.get('template_id') is not None)

        _templates_data_local = []
        for template in app_state.all_available_templates:
            if template.get('id') in template_ids_in_selected_hub_documents:
                _templates_data_local.append(template)

        app_state.template_titles_for_combo = [t.get('title', f"Untitled Template (ID: {t.get('id')})") for t in
                                               _templates_data_local]

        if app_state.template_titles_for_combo:
            dpg.configure_item("template_combo", items=app_state.template_titles_for_combo, enabled=True)
            log_message(
                f"✅ Found {len(app_state.template_titles_for_combo)} templates associated with Hub {app_state.selected_hub_id}.")
        else:
            dpg.configure_item("template_combo", items=[], enabled=False)
            log_message(f"❌ No templates found for documents within Hub {app_state.selected_hub_id}.")

        dpg.configure_item("status_text",
                           default_value=f"✅ Hub {app_state.selected_hub_id} loaded. Select folder and template.")

    def folder_selected_callback():
        selected_folder_value = dpg.get_value("folder_combo")
        if not selected_folder_value: return

        if "Root of Hub (ID: " in selected_folder_value:
            root_hub_id_str = selected_folder_value.split(" (ID: ")[1].rstrip(")")
            app_state.set_selected_folder(int(root_hub_id_str), selected_folder_value)
            log_message(
                f"✅ Selected Root of Hub. Uploads will go to Hub ID: {app_state.selected_folder_id} (acting as parent folder).")
        elif "(Folder ID: " in selected_folder_value:
            folder_id_str = selected_folder_value.split(" (Folder ID: ")[1].rstrip(")")
            app_state.set_selected_folder(int(folder_id_str), selected_folder_value)
            log_message(f"✅ Selected Folder '{selected_folder_value}'. Folder ID: {app_state.selected_folder_id}")
        else:
            app_state.set_selected_folder(None, selected_folder_value)
            log_message(f"⚠️ Could not parse folder ID from '{selected_folder_value}'. Uploads might fail.")

        if dpg.does_item_exist("upload_button"):
            if app_state.selected_folder_id is not None and app_state.selected_template_id is not None and _selected_file_path:
                dpg.configure_item("upload_button", enabled=True)
            else:
                dpg.configure_item("upload_button", enabled=False)

    def template_selected_callback():
        selected_template_title = dpg.get_value("template_combo")

        _selected_template_details_local = next(
            (t for t in app_state.all_available_templates if t.get('title') == selected_template_title), None)

        if _selected_template_details_local:
            app_state.set_selected_template(_selected_template_details_local.get('id'),
                                            _selected_template_details_local)
            print(f"DEBUG: Selected Template Details for '{selected_template_title}':")
            print(json.dumps(app_state.selected_template_details, indent=2))

            log_message(
                f"✅ Template '{selected_template_title}' (ID: {app_state.selected_template_id}) selected for use.")
            if dpg.does_item_exist("generate_template_button"):
                dpg.configure_item("generate_template_button", enabled=True)
        else:
            log_message(f"⚠️ Could not find details for selected template '{selected_template_title}'.")
            if dpg.does_item_exist("generate_template_button"):
                dpg.configure_item("generate_template_button", enabled=False)
            app_state.set_selected_template(None, None)

        if dpg.does_item_exist("upload_button"):
            if app_state.selected_folder_id is not None and app_state.selected_template_id is not None and _selected_file_path:
                dpg.configure_item("upload_button", enabled=True)
            else:
                dpg.configure_item("upload_button", enabled=False)

    def generate_template_callback():
        if not app_state.selected_template_details:
            log_message("❌ No template selected to generate Excel from.")
            return

        template_title = app_state.selected_template_details.get('title', 'Untitled_Template')
        template_id = app_state.selected_template_details.get('id', 'unknown_id')
        output_filename = f"{template_title}_template_ID_{template_id}.xlsx"

        headers = ["Title", "Description"]

        field_validation_info = {}

        if 'fields' in app_state.selected_template_details and app_state.selected_template_details['fields']:
            for field in app_state.selected_template_details['fields']:
                field_name = field.get('name_singular') or field.get('name_plural') or f"Field ID: {field.get('id')}"
                if field_name not in headers:
                    headers.append(field_name)

                field_validation_info[field_name] = {
                    "field_type": field.get('field_type'),
                    "options": field.get('options')
                }
        elif 'fields' in app_state.selected_template_details and not app_state.selected_template_details['fields']:
            log_message(
                f"⚠️ Selected template '{template_title}' has no custom fields defined. Generating Excel file with 'Title' and 'Description' only.")
        else:
            log_message(
                f"⚠️ Selected template '{template_title}' details do not contain a 'fields' array. Generating Excel file with 'Title' and 'Description' only.")

        try:
            df = pd.DataFrame(columns=headers)

            with pd.ExcelWriter(output_filename, engine='openpyxl') as writer:
                df.to_excel(writer, index=False, sheet_name='Data Input')
                workbook = writer.book
                sheet = writer.sheets['Data Input']

                for col_idx, header_name in enumerate(headers):
                    if header_name == "Title" or header_name == "Description":
                        continue

                    if header_name in field_validation_info:
                        field_info = field_validation_info[header_name]
                        if field_info['field_type'] == 'PICKER' and field_info['options']:
                            options_list = [option.get('title') for option in field_info['options'] if
                                            isinstance(option, dict) and 'title' in option]
                            options_str = ','.join(options_list)

                            col_letter = chr(65 + col_idx)
                            validation_range = f"{col_letter}2:{col_letter}1000"

                            dv = DataValidation(type="list", formula1=f'"{options_str}"', allow_blank=True)
                            dv.error = 'Value must be selected from the dropdown list.'
                            dv.errorTitle = 'Invalid Entry'
                            dv.showErrorMessage = True
                            dv.showDropDown = True

                            sheet.add_data_validation(dv)
                            dv.add(validation_range)
                            log_message(f"Applied dropdown validation to '{header_name}' with options: {options_str}")

            log_message(f"✅ Excel template generated: {output_filename}")
            dpg.set_value("upload_file_path", output_filename)
            globals()['_selected_file_path'] = output_filename
            if app_state.selected_folder_id is not None and app_state.selected_template_id is not None and globals()[
                '_selected_file_path']:
                dpg.configure_item("upload_button", enabled=True)

        except Exception as e:
            log_message(f"❌ Error generating Excel template: {e}")
            import traceback
            traceback.print_exc()

    def _perform_duplicate_check(excel_df: pd.DataFrame) -> list:
        duplicate_details = []
        if app_state.selected_hub_id is None or app_state.selected_folder_id is None:
            log_message("❌ Cannot perform duplicate check: Hub or Folder not selected.")
            return []

        existing_docs_in_target_location = {}
        for doc in app_state.all_documents:
            if doc.get('document_hub_id') == app_state.selected_hub_id and \
                    doc.get('parent_folder_id') == app_state.selected_folder_id and \
                    doc.get('title'):
                existing_docs_in_target_location[doc['title'].lower()] = doc

        for index, row in excel_df.iterrows():
            incoming_title = str(row.get("Title", "")).strip()
            incoming_description = str(row.get("Description", "")).strip()

            if not incoming_title:
                log_message(f"⚠️ Warning: Excel row {index + 2} has no title. Skipping for duplicate check.")
                continue

            if incoming_title.lower() in existing_docs_in_target_location:
                existing_doc = existing_docs_in_target_location[incoming_title.lower()]
                duplicate_details.append({
                    "incoming_title": incoming_title,
                    "incoming_description": incoming_description,
                    "existing_title": existing_doc.get('title', 'N/A'),
                    "existing_description": existing_doc.get('description', 'N/A'),
                    "existing_id": existing_doc.get('id', 'N/A'),
                    "excel_row_index": index
                })

        return duplicate_details

    def _duplicate_dialog_callback(sender, app_data, user_data):
        dpg.delete_item("duplicate_check_modal")

        action = user_data.get("action")

        df_to_upload = user_data['df_to_upload'].copy()  # Use passed df_to_upload

        if action == "skip_duplicates":
            duplicate_excel_row_indices = [dup['excel_row_index'] for dup in user_data['duplicate_details']]
            df_to_upload = df_to_upload.drop(duplicate_excel_row_indices).reset_index(drop=True)
            log_message(f"Skipping {len(user_data['duplicate_details'])} duplicate document(s).")
            if df_to_upload.empty:
                log_message("No non-duplicate documents remaining for upload. Upload cancelled.")
                dpg.configure_item("upload_button", enabled=True)
                return
        elif action == "upload_all_anyway":
            log_message("Proceeding with upload, including duplicates.")
        elif action == "cancel_upload":
            log_message("Upload cancelled by user.")
            dpg.configure_item("upload_button", enabled=True)
            return

        log_message(f"🚀 Starting upload process for: {user_data['file_path']}")  # Use passed file_path
        dpg.configure_item("upload_button", enabled=False)

        upload_manager.upload_documents_from_excel(
            _main_window_config_global,  # Use global config
            df_to_upload,
            app_state.selected_hub_id,
            app_state.selected_folder_id,
            app_state.selected_template_details,
            log_callback=log_message,
            on_success_callback=fetch_hubs_callback
        )
        dpg.configure_item("upload_button", enabled=True)
        log_message("Upload process finished (check logs for details).")

    def file_selected_callback(sender, app_data):
        global _selected_file_path
        if app_data and app_data.get('file_path_name'):
            _selected_file_path = app_data['file_path_name']
            dpg.set_value("upload_file_path", _selected_file_path)
            log_message(f"✅ File selected for upload: {_selected_file_path}")
        else:
            _selected_file_path = ""
            dpg.set_value("upload_file_path", "")
            log_message("❌ No file selected.")

        if app_state.selected_folder_id is not None and app_state.selected_template_id is not None and _selected_file_path:
            dpg.configure_item("upload_button", enabled=True)
        else:
            dpg.configure_item("upload_button", enabled=False)

    def upload_documents_callback():
        global _current_upload_df

        if not _selected_file_path:
            log_message("❌ No Excel file selected for upload.")
            return
        if app_state.selected_hub_id is None:
            log_message("❌ No Document Hub selected.")
            return
        if app_state.selected_folder_id is None:
            log_message("❌ No Folder selected.")
            return
        if app_state.selected_template_details is None:
            log_message("❌ No Template selected.")
            return

        try:
            _current_upload_df = pd.read_excel(_selected_file_path)
            _current_upload_df._file_path = _selected_file_path
            if _current_upload_df.empty:
                log_message("❌ Selected Excel file is empty. No documents to upload.")
                return
        except Exception as e:
            log_message(f"❌ Error reading Excel file '{_selected_file_path}': {e}")
            import traceback
            traceback.print_exc()
            return

        duplicate_details = _perform_duplicate_check(_current_upload_df)

        if duplicate_details:
            with dpg.window(label="Duplicate Documents Found", modal=True, show=True, tag="duplicate_check_modal",
                            width=650, height=350, no_close=True):
                dpg.add_text("The following document titles already exist in the target location:", wrap=0)
                dpg.add_spacer(height=5)

                with dpg.table(header_row=True, resizable=True, policy=dpg.mvTable_SizingStretchProp,
                               borders_outerH=True, borders_innerV=True, borders_innerH=True, borders_outerV=True):
                    dpg.add_table_column(label="Incoming Title", width_fixed=True, init_width_or_weight=150)
                    dpg.add_table_column(label="Incoming Desc", width_fixed=True, init_width_or_weight=150)
                    dpg.add_table_column(label="Existing Title", width_fixed=True, init_width_or_weight=150)
                    dpg.add_table_column(label="Existing Desc", width_fixed=True, init_width_or_weight=150)

                    for dup in duplicate_details:
                        with dpg.table_row():
                            dpg.add_text(f"{str(dup.get('incoming_title', 'N/A'))[:25]}...", wrap=False)
                            dpg.add_text(f"{str(dup.get('incoming_description', 'N/A'))[:25]}...", wrap=False)
                            dpg.add_text(f"{str(dup.get('existing_title', 'N/A'))[:25]}...")
                            dpg.add_text(f"{str(dup.get('existing_description', 'N/A'))[:25]}...", wrap=False)

                dpg.add_spacer(height=10)
                dpg.add_text("How would you like to proceed?", wrap=0)

                with dpg.group(horizontal=True):
                    dpg.add_button(label="Skip Duplicates", callback=_duplicate_dialog_callback,
                                   user_data={"action": "skip_duplicates", "duplicate_details": duplicate_details,
                                              "df_to_upload": _current_upload_df,
                                              "file_path": _selected_file_path})  # Pass df_to_upload and file_path
                    dpg.add_button(label="Upload All Anyway", callback=_duplicate_dialog_callback,
                                   user_data={"action": "upload_all_anyway", "df_to_upload": _current_upload_df,
                                              "file_path": _selected_file_path})
                    dpg.add_button(label="Cancel Upload", callback=_duplicate_dialog_callback,
                                   user_data={"action": "cancel_upload"})
            log_message(f"Found {len(duplicate_details)} potential duplicate(s). Waiting for user input.")
        else:
            log_message("No duplicate documents found. Proceeding with upload.")
            upload_manager.upload_documents_from_excel(
                _main_window_config_global,
                _current_upload_df,
                app_state.selected_hub_id,
                app_state.selected_folder_id,
                app_state.selected_template_details,
                log_callback=log_message,
                on_success_callback=fetch_hubs_callback
            )
            dpg.configure_item("upload_button", enabled=True)
            log_message("Upload process finished (check logs for details).")

    # Callback for opening the misc tools window
    def open_misc_tools_callback():
        dpg.show_item("misc_tools_window")
        # Trigger refresh of its internal combos when shown
        # FIX APPLIED HERE
        misc_tools_window.refresh_misc_data_and_combos(None, None)

    with dpg.window(label="Main Window", tag="main_window", show=False):
        # Viewport Menu Bar for settings and tools
        with dpg.viewport_menu_bar():
            with dpg.menu(label="File"):
                dpg.add_menu_item(label="Settings", callback=show_config)
            with dpg.menu(label="Tools"):
                dpg.add_menu_item(label="Miscellaneous Tools", callback=_show_and_refresh_misc_tools)

        dpg.add_separator()

        with dpg.group(horizontal=True):
            dpg.add_text(status_message, tag="status_text")

        dpg.add_separator()

        with dpg.collapsing_header(label="Step 1: Select Target Location", default_open=True):
            dpg.add_checkbox(label="Force API Refresh (Bypass Cache)", tag="force_refresh_checkbox",
                             default_value=False)
            dpg.add_button(label="Fetch Document Hubs & Templates", tag="fetch_button", callback=fetch_hubs_callback)
            dpg.add_combo(label="Document Hub", tag="hub_combo", enabled=False, callback=hub_selected_callback,
                          width=300)
            dpg.add_combo(label="Folder", tag="folder_combo", enabled=False, callback=folder_selected_callback,
                          width=300)
            dpg.add_combo(label="Template", tag="template_combo", enabled=False, callback=template_selected_callback,
                          width=300)

        with dpg.collapsing_header(label="Step 2: Generate Excel Template", default_open=True):
            dpg.add_button(label="Generate Excel Template", tag="generate_template_button", enabled=False,
                           callback=generate_template_callback)

        with dpg.collapsing_header(label="Step 3: Select File and Upload", default_open=True):
            dpg.add_text("Upload Excel file generated from template.")
            dpg.add_input_text(label="File Path", tag="upload_file_path", default_value="", width=400, enabled=False)
            dpg.add_button(label="Browse File...", callback=lambda: dpg.show_item("file_dialog_id"))
            dpg.add_button(label="Upload Selected File", tag="upload_button", enabled=False,
                           callback=upload_documents_callback)

        dpg.add_separator()

        # This button is now redundant if using the menu bar, but we'll leave it and its callback fixed
        dpg.add_button(label="Open Other Tools", callback=open_misc_tools_callback)

        with dpg.file_dialog(directory_selector=False, show=False, callback=file_selected_callback, id="file_dialog_id",
                             width=700, height=400):
            dpg.add_file_extension(".xlsx", custom_text="Excel Files (*.xlsx)")
            dpg.add_file_extension(".*", custom_text="All Files")
        dpg.hide_item("file_dialog_id")